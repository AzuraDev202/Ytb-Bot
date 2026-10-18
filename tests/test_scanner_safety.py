import datetime as dt
import inspect
from pathlib import Path
import queue
import sys
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from ytb_bot.core.engine import ScanEngine
from ytb_bot.cli import scanner
from ytb_bot.legacy_ui.desktop import YtbBotApp


def youtube_fixture():
    youtube = MagicMock()
    youtube.search.return_value.list.return_value.execute.return_value = {
        'items': [{'snippet': {'channelId': 'creator', 'title': '=1+2'}}]}
    youtube.channels.return_value.list.return_value.execute.return_value = {'items': [
        {'id': 'creator', 'snippet': {'title': '=1+1', 'description': 'hello@example.com',
         'publishedAt': dt.datetime.now(dt.timezone.utc).isoformat()},
         'statistics': {'viewCount': '1000'}}]}
    return youtube


class ScannerSafetyTests(unittest.TestCase):
    def test_restart_rejected_until_previous_worker_exits(self):
        paused, resume = threading.Event(), threading.Event()
        source, first_line = inspect.getsourcelines(ScanEngine._run)
        pause_line = first_line + next(i for i, line in enumerate(source) if line.strip() == 'finally:') + 1
        class PausingEngine(ScanEngine):
            def _run(self, *args):
                def trace(frame, event, arg):
                    if frame.f_code is ScanEngine._run.__code__ and event == 'line' and frame.f_lineno == pause_line:
                        paused.set()
                        resume.wait(5)
                    return trace
                sys.settrace(trace)
                try:
                    super()._run(*args)
                finally:
                    sys.settrace(None)
        with tempfile.TemporaryDirectory() as directory:
            youtube = youtube_fixture()
            engine = PausingEngine(directory, lambda _: youtube)
            engine.start('key', ['one'])
            self.assertTrue(paused.wait(5))
            worker = engine._thread
            try:
                with self.assertRaises(ValueError):
                    engine.start('key', ['two'])
            finally:
                resume.set()
                worker.join(5)
                engine.wait()
            self.assertEqual(len(engine._store.runs()), 1)
            self.assertEqual(engine._store.result()['status'], 'complete')

    def test_cli_exports_public_titles_as_text(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'result.xlsx'
            with patch.object(scanner, 'get_youtube_client', return_value=youtube_fixture()), patch.object(scanner, 'load_keywords_from_excel', return_value=['one']), patch('builtins.print'):
                scanner.scan_channels_from_excel(output_excel=output)
            self.assert_titles_are_text(output)

    def assert_titles_are_text(self, path):
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            columns = {cell.value: cell.column for cell in sheet[1]}
            for column, value in [('Channel Name', '=1+1'), ('Latest Video', '=1+2')]:
                cell = sheet.cell(2, columns[column])
                self.assertEqual(cell.data_type, 's')
                self.assertEqual(cell.value, value)
        finally:
            workbook.close()

    def test_legacy_keeps_matches_on_later_network_failure(self):
        youtube = youtube_fixture()
        first = youtube.search.return_value.list.return_value.execute.return_value
        youtube.search.return_value.list.return_value.execute.side_effect = [first, OSError('offline')]
        callbacks = ['_scan_match_count', '_scan_progress', '_show_excel', '_scan_notice', '_scan_failed', '_scan_finished']
        app = SimpleNamespace(stop_requested=threading.Event(), ui_events=queue.Queue(), log=MagicMock(),
                              **{name: MagicMock() for name in callbacks})
        with tempfile.TemporaryDirectory() as directory:
            with patch('ytb_bot.legacy_ui.desktop.APP_DIR', Path(directory)), patch('ytb_bot.legacy_ui.desktop.build', return_value=youtube):
                YtbBotApp.run_scanner(app, 'key', ['one', 'two'], 90, 300, 1)
            output = Path(directory) / 'channel_results.xlsx'
            self.assertTrue(output.exists(), 'Accepted matches must survive a later API/network error')
            self.assert_titles_are_text(output)
            self.assertTrue(any(callback is app._scan_failed for callback, _ in list(app.ui_events.queue)))
